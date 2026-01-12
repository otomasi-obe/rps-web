#!/usr/bin/env python3
"""Analyze RPS.docx template structure"""

from docx import Document
import json

doc = Document('RPS.docx')

print('=== DOCX ANALYSIS ===')
print(f'Total tables: {len(doc.tables)}')
print(f'Total paragraphs: {len(doc.paragraphs)}')

# Analisis setiap tabel
for table_idx, table in enumerate(doc.tables):
    print(f'\n=== TABLE {table_idx} ===')
    print(f'Rows: {len(table.rows)}, Cols: {len(table.columns)}')
    print('Content:')
    for row_idx, row in enumerate(table.rows[:20]):  # Show first 20 rows
        row_content = []
        for col_idx, cell in enumerate(row.cells):
            text = cell.text.strip()
            if text:
                row_content.append(f'Col{col_idx}: {text[:100]}')
        if row_content:
            print(f'  Row {row_idx}: {" | ".join(row_content)}')
        elif row_idx < 5:
            print(f'  Row {row_idx}: (empty)')
